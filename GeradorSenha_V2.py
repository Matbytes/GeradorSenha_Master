#Por Matheus Silva
#Projeto iniciado 28/8/2026

#Gerador de Senha V2
#Bibliotecas importantes
import random
import openpyxl
from openpyxl import load_workbook

wb = openpyxl.load_workbook("senhas.xlsx")
sheet = wb.active
sheet.title = "Senhas e serviços"

varControlD = True #Variavel booleana para controle do loop
varContro = 6 #Variavel de controle para o loop principal do código
linhaIni = 1

#Loop para verificar se há linhas ocupadas ou desocupadas 
while varControlD:
    celNum = sheet.cell(row=linhaIni,column=3)
    celNumD = celNum.data_type

    if celNumD == "s":
        linhaIni += 1
    else:
        global linhaDese
        linhaDese = linhaIni 
        break

#Loop para a geração de senhas
while varContro < 7:

    print("Gerador de Senha v2\n")

    
    servicoNome = input("Nome do serviço utilizado: ")
    nomeUsu = input("Nome de usuario: ")
    numCarac = input("Numero de caracteres da senha: ")
    numCarac = int(numCarac)
    
    #Loop para estabelecer um limite m inimo de quatro caracteres para a senha
    while numCarac < 4:
        print("Erro: A senha deve ter no minimo 4 caracteres")
        numCarac = input("Numero de caracteres da senha: ")
        numCarac = int(numCarac)
        if numCarac >= 4:
            break
    
    querNum = input("Deseja ter numeros? S/N: ")
    querEspe = input("Deseja ter simbolos? S/N: " )
    querLetrasMais = input("Deseja ter letras maisculas? S/N: ")
    querLetrasMin = input("Deseja ter letras minusculas? S/N: ")
    
    letrasMa = ["A","B","C","D","E","F","G", "H","I","J","K","L","M","N","O","P","Q","R","S","T","U","V","W","X","Y","Z"]
    letrasmi = ["a","b","c","d","e","f","g","h","i","j","k","l","m","n","o","p","q","r","s","t","u","v","w","x","y","z"]
    numeros = ["0","1","2","3","4","5","6","7","8","9"]
    caracEspeciais = ["!","@","#","$","&","*"]
    universo = []
    senha = []


    f = 3
    x = random.randint(0, 25)
    b = random.randint(0,25)
    y = random.randint(0,9)
    z = random.randint(0,5)
    w = random.randint(0,f)

    Listas = [letrasMa[x], letrasmi[b], numeros[y], caracEspeciais[z]]


    senha.append(letrasMa[x])
    senha.append(letrasmi[b])
    senha.append(numeros[y])
    senha.append(caracEspeciais[z])

    random.shuffle(senha)
    Senha_Final = "".join(senha)

    #Nessa condição vai ocorrer a expansão do numero de caracteres definido pelo usuario
    if numCarac > 4:
        numCarac -= 4
        universo.extend(letrasMa)
        universo.extend(letrasmi)
        universo.extend(numeros)
        universo.extend(caracEspeciais)
        for i in range(0,numCarac):
            global t
            t = random.choice(universo)
            if t in senha:
                t = random.choice(universo)
                senha.append(t)
            else:
                senha.append(t)
        random.shuffle(senha)
    else:
        pass
        
    match querNum:
        case "S":
            pass
        case "s":
            pass
        case "N":
            senha.remove(numeros[y])
            Listas.remove(numeros[y])
            for t in numeros:
                if t in senha:
                    senha.remove(t)
            f -= 1
            Listas.append(Listas[w])
            senha.append(Listas[w])
        case "n":
            senha.remove(numeros[y])
            Listas.remove(numeros[y])
            for t in numeros:
                if t in senha:
                    senha.remove(t)
            f -= 1
            Listas.append(Listas[w])
            senha.append(Listas[w])

    match querEspe:
        case "S":
            pass
        case "s":
            pass
        case "N":
            senha.remove(caracEspeciais[z])
            Listas.remove(caracEspeciais[z])
            for t in caracEspeciais:
                if t in senha:
                    senha.remove(t)
            f -= 1
            Listas.append(Listas[w])
            senha.append(Listas[w])
        case "n":
            senha.remove(caracEspeciais[z])
            Listas.remove(caracEspeciais[z])
            for t in caracEspeciais:
                if t in senha:
                    senha.remove(t)
            f -= 1
            Listas.append(Listas[w])
            senha.append(Listas[w])

    match querLetrasMais:
        case "S":
            pass
        case "s":
            pass
        case "N":
            senha.remove(letrasMa[x])
            Listas.remove(letrasMa[x])
            for t in letrasMa:
                if t in senha:
                    senha.remove(t)
            f -= 1
            Listas.append(Listas[w])
            senha.append(Listas[w])
        case "n":
            senha.remove(letrasMa[x])
            Listas.remove(letrasMa[x])
            for t in letrasMa:
                if t in senha:
                    senha.remove(t)
            f -= 1
            Listas.append(Listas[w])
            senha.append(Listas[w])

    match querLetrasMin:
        case "S":
            pass
        case "s":
            pass
        case "N":
            senha.remove(letrasmi[b])
            Listas.remove(letrasmi[b])
            for t in letrasmi:
                if t in senha:
                    senha.remove(t)
            f -= 1
            Listas.append(Listas[w])
            senha.append(Listas[w])
        case "n":
            senha.remove(letrasmi[b])
            Listas.remove(letrasmi[b])
            for t in letrasmi:
                if t in senha:
                    senha.remove(t)
            f -= 1
            Listas.append(Listas[w])
            senha.append(Listas[w])
    

    if f == 0:
        print("Erro: Selecione ao menos uma categoria de caracteres!")
        exit()

    def exit():
        return

    #Dando uma última embaralhada na senha
    random.shuffle(senha)
    
    #Formata a lista senha para uma fomra mais "limpa" 
    Senha_Final = "".join(senha)

    print("\nServiço usado: ", servicoNome)
    print("Nome de Usuario: ",nomeUsu)
    print("Senha: ",Senha_Final)
    
    #Aqui inicia o preenchimento na planilha do Excel
    sheet["A1"] = "Serviço"
    sheet["B1"] = "Usuario"
    sheet["C1"] = "Senha"

    serCel = sheet.cell(row=linhaIni, column=1)
    usuCel = sheet.cell(row=linhaIni, column=2)
    seCel = sheet.cell(row=linhaIni, column=3)

    serCel.value = servicoNome
    usuCel.value = nomeUsu
    seCel.value = Senha_Final

    serData = serCel.data_type
    usuData = usuCel.data_type
    seData = seCel.data_type

    #Se a celula atual estiver ocupada, então o valor para a próxima célula 
    match seData:
        case "n":
            print("Vazio")
        case "s":
            linhaIni += 1

    wb.save("senhas.xlsx")

    try:
        input("\nPressiona enter para continuar ou CTRL + c para fechar o programa\n")
    except KeyboardInterrupt:
        print("\n Blz, saindo..")
        varContro = 8
    else:
        pass
    